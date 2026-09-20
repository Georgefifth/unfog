#!/usr/bin/env python3
"""Build Unfog pitch deck (pptx) from gallery screenshots."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import os

HERE = os.path.dirname(os.path.abspath(__file__))
GAL = os.path.join(HERE, "gallery")
OUT = os.path.join(HERE, "unfog-deck.pptx")

INK = RGBColor(0x2A, 0x27, 0x3F)
MUTED = RGBColor(0x6E, 0x6A, 0x83)
ACCENT = RGBColor(0x4F, 0x46, 0xE5)
PAPER = RGBColor(0xFA, 0xF9, 0xF6)
RED = RGBColor(0xB9, 0x1C, 0x1C)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height


def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(1, 0, 0, SW, SH)  # rect
    bg.fill.solid(); bg.fill.fore_color.rgb = PAPER
    bg.line.fill.background(); bg.shadow.inherit = False
    return s


def text(s, x, y, w, h, runs, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.15):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    items = runs if isinstance(runs, list) else [(runs, {})]
    for i, (txt, opts) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = opts.get("align", align)
        p.line_spacing = spacing
        p.space_after = Pt(opts.get("space", 6))
        r = p.add_run(); r.text = txt
        f = r.font
        f.size = Pt(opts.get("size", size))
        f.bold = opts.get("bold", bold)
        f.color.rgb = opts.get("color", color)
        f.name = "Segoe UI"
    return tb


def pic(s, path, x, y, w=None, h=None):
    return s.shapes.add_picture(path, x, y, w, h)


def bullet(s, x, y, w, h, items, size=17):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    for i, (head, sub) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.15; p.space_after = Pt(4)
        r = p.add_run(); r.text = head
        r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = INK; r.font.name = "Segoe UI"
        if sub:
            p2 = tf.add_paragraph()
            p2.line_spacing = 1.1; p2.space_after = Pt(10)
            r2 = p2.add_run(); r2.text = sub
            r2.font.size = Pt(size - 3); r2.font.color.rgb = MUTED; r2.font.name = "Segoe UI"
    return tb


def kicker(s, txt):
    text(s, Inches(0.7), Inches(0.45), Inches(8), Inches(0.4), txt, size=13, bold=True, color=ACCENT)


def footer(s, txt="github.com/Georgefifth/unfog · unfog-oaye.onrender.com"):
    text(s, Inches(0.7), SH - Inches(0.55), Inches(11.9), Inches(0.35), txt, size=11, color=MUTED)


# ── 1. Title ──────────────────────────────────────────────
s = slide()
text(s, Inches(0.9), Inches(1.7), Inches(11.5), Inches(1), "BUNNIEX HACKATHON", size=16, bold=True, color=ACCENT)
text(s, Inches(0.9), Inches(2.1), Inches(11.5), Inches(1.6), "Unfog", size=96, bold=True)
text(s, Inches(0.9), Inches(3.6), Inches(11.5), Inches(1), "Snap confusing paperwork → get plain words, real deadlines, and a plan.", size=26, color=MUTED)
text(s, Inches(0.9), Inches(5.6), Inches(11.5), Inches(1),
     "Powered by Featherless AI · Qwen3-VL + Qwen2.5-72B\nLive: unfog-oaye.onrender.com   ·   Code: github.com/Georgefifth/unfog", size=15, color=MUTED)

# ── 2. Problem ────────────────────────────────────────────
s = slide()
kicker(s, "THE PROBLEM")
text(s, Inches(0.7), Inches(0.85), Inches(12), Inches(1.2), "Paperwork is written to confuse.", size=40, bold=True)
bullet(s, Inches(0.9), Inches(2.2), Inches(6.2), Inches(4.5), [
    ("🧾  Medical bills", "CPT codes, 'adjustments', insurance math nobody can verify"),
    ("🅿️  Tickets & fines", "escalating penalties buried in legalese"),
    ("🏠  Lease notices", "auto-renewal traps, quiet fee changes"),
    ("🛡️  Insurance denials", "written to discourage appeals"),
    ("🌍  And the burden is unequal", "hardest for students, immigrants, the elderly — anyone reading in a second language"),
], size=17)
box = s.shapes.add_shape(1, Inches(7.6), Inches(2.3), Inches(4.9), Inches(3.4))
box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0xFF, 0xFB, 0xEB); box.line.color.rgb = RGBColor(0xF5, 0x9E, 0x0B); box.line.width = Pt(1.5)
tf = box.text_frame; tf.word_wrap = True; tf.margin_left = tf.margin_right = Inches(0.3); tf.margin_top = Inches(0.25)
p = tf.paragraphs[0]; r = p.add_run(); r.text = "The cost isn't confusion."
r.font.size = Pt(20); r.font.bold = True; r.font.color.rgb = INK; r.font.name = "Segoe UI"
p2 = tf.add_paragraph(); p2.space_before = Pt(10)
r2 = p2.add_run(); r2.text = "It's missed deadlines. Overpaid fees. Appeals never filed. Money lost to fine print."
r2.font.size = Pt(18); r2.font.color.rgb = INK; r2.font.name = "Segoe UI"
footer(s)

# ── 3. Solution (screenshot) ─────────────────────────────
s = slide()
kicker(s, "THE SOLUTION")
text(s, Inches(0.7), Inches(0.85), Inches(12), Inches(1), "Photo in → clarity and a plan out", size=40, bold=True)
pic(s, f"{GAL}/02-result-top.png", Inches(0.7), Inches(1.9), w=Inches(11.9))
footer(s)

# ── 4. What it does ──────────────────────────────────────
s = slide()
kicker(s, "WHAT IT DOES")
text(s, Inches(0.7), Inches(0.85), Inches(12), Inches(1), "Not a summarizer — a decoder", size=40, bold=True)
bullet(s, Inches(0.9), Inches(2.0), Inches(5.9), Inches(5), [
    ("📖 Plain-language summary + ELI5", "what it says, in words anyone understands"),
    ("🎯 Grounded facts", "hover a fact → see where it sits on the page (VL bounding boxes)"),
    ("⏰ Deadlines", "countdowns + one-tap calendar export"),
    ("🛡️ Scam check", "is this document even real? gift-card demands, arrest threats flagged"),
], size=16)
bullet(s, Inches(7.0), Inches(2.0), Inches(5.6), Inches(5), [
    ("✅ Action checklist", "every step has a 'how do I do this' button"),
    ("✍️ Drafts that know your facts", "dispute letters pre-filled with real account numbers"),
    ("🎭 Call rehearsal", "practice the scary phone call — by voice — vs an AI rep"),
    ("💬 Grounded Q&A + 🌍 12 languages", "answers quote your document, not generic advice"),
], size=16)
footer(s)

# ── 5. Scam check (screenshot) ───────────────────────────
s = slide()
kicker(s, "TRUST LAYER")
text(s, Inches(0.7), Inches(0.85), Inches(12), Inches(1), "\"Is this even real?\"", size=40, bold=True)
pic(s, f"{GAL}/07-scam-check.png", Inches(0.7), Inches(1.9), w=Inches(11.9))
footer(s)

# ── 6. How it works ──────────────────────────────────────
s = slide()
kicker(s, "HOW IT'S BUILT")
text(s, Inches(0.7), Inches(0.85), Inches(12), Inches(1), "Featherless AI, end to end", size=40, bold=True)
steps = [("📷", "Photo / PDF / text"), ("🖼️", "Pillow + PyMuPDF\npreprocess"), ("👁️", "Qwen3-VL-30B\nvision → JSON + bbox"),
         ("✍️", "Qwen2.5-72B\ndrafts · chat · roleplay"), ("📱", "Vanilla JS UI\nSSE streaming")]
x = Inches(0.7); w = Inches(2.25); gap = Inches(0.25)
for i, (ico, label) in enumerate(steps):
    b = s.shapes.add_shape(1, x, Inches(2.3), w, Inches(1.7))
    b.fill.solid(); b.fill.fore_color.rgb = RGBColor(0xEF, 0xEE, 0xFC); b.line.fill.background()
    tf = b.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = ico; r.font.size = Pt(28)
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run(); r2.text = label; r2.font.size = Pt(13); r2.font.color.rgb = INK; r2.font.name = "Segoe UI"
    x += w + gap
bullet(s, Inches(0.9), Inches(4.5), Inches(11.5), Inches(2.2), [
    ("Resilience: multi-model fallback chain", "gated models, capacity exhaustion, and '!'-token floods rerouted mid-flight"),
    ("Streaming everything", "SSE token streaming + elapsed timers make 25s cold starts feel instant"),
    ("Privacy by default", "no server storage — history lives in the browser's localStorage"),
], size=15)
footer(s)

# ── 7. Challenges ────────────────────────────────────────
s = slide()
kicker(s, "CHALLENGES → FIXES")
text(s, Inches(0.7), Inches(0.85), Inches(12), Inches(1), "Open models need babysitting", size=40, bold=True)
rows = [
    ("Pure '!'-token floods from upstream", "detect + reroute to next model mid-stream"),
    ("Llama weights gated, VL capacity exhausted", "fallback chain: Qwen2.5-72B → Qwen3-32B → Mistral-Small"),
    ("Models wrap JSON in chatter", "defensive extraction + retry loop"),
    ("Blocking HTTP froze the event loop", "inference moved to threadpool"),
    ("25s+ cold starts", "SSE streaming, elapsed timers, honest placeholders"),
]
y = 2.1
for prob, fix in rows:
    text(s, Inches(0.9), Inches(y), Inches(6.3), Inches(0.7), "✗ " + prob, size=15, color=RED)
    text(s, Inches(7.3), Inches(y), Inches(5.4), Inches(0.7), "→ " + fix, size=15, color=INK)
    y += 0.85
footer(s)

# ── 8. Close ─────────────────────────────────────────────
s = slide()
text(s, Inches(0.9), Inches(1.9), Inches(11.5), Inches(1.2), "Nobody should lose money to fine print.", size=44, bold=True)
text(s, Inches(0.9), Inches(3.3), Inches(11.5), Inches(1.8),
     "▶  Live demo: unfog-oaye.onrender.com\n📦  Code: github.com/Georgefifth/unfog\n🎬  Video: demo.mp4 (narrated, 56s)",
     size=22, color=INK)
text(s, Inches(0.9), Inches(5.6), Inches(11.5), Inches(1),
     "Next: verified-org lookup · document diffing · deadline push alerts · PWA",
     size=15, color=MUTED)

prs.save(OUT)
print("saved", OUT)
